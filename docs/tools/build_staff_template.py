"""Builds docs/templates/ABF-FST_Staff_Accounts_Template.xlsx -- the sheet the
PI fills in for manage.py create_staff_accounts. Run with the backend venv:
    backend/.venv/Scripts/python.exe docs/tools/build_staff_template.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parents[1] / "templates" / "ABF-FST_Staff_Accounts_Template.xlsx"
ROLES = [
    ("Contact RA", "Contacts assigned organisations, records respondents, sends invitations and reminders"),
    ("Field / Digital Coordinator", "Runs day-to-day fieldwork; assigns and verifies cases; replacements and withdrawals"),
    ("QUAN/Kobo QA RA", "Checks questionnaire quality: QA Queue and QA Exceptions"),
    ("KII RA", "Arranges, runs and processes Key Informant Interviews"),
    ("Documentary RA", "Collects, assesses and codes documentary evidence"),
    ("Data Analyst", "Dashboards, Reports and the de-identified analysis export (read-only)"),
    ("Supervisor (read-only)", "Read-only oversight of almost every screen"),
    ("PI / System Admin", "Full access, including the audit log and all exports"),
]
# Main-400 cases per province on the live system, 15 Sep 2026 -- to plan how to split them.
CASES = [("Harare", 252), ("Bulawayo", 37), ("Manicaland", 25), ("Mashonaland West", 22), ("Mashonaland East", 20),
         ("Midlands", 20), ("Masvingo", 10), ("Mashonaland Central", 9), ("Matabeleland South", 3), ("Matabeleland North", 2)]
HEADER = ["Full name", "Email", "Role", "Provinces (Contact RA only)", "Username (optional)"]

NAVY, GOLD, INPUT = "1B3350", "A67C27", "FFF7CC"
font = lambda **kw: Font(name="Arial", size=kw.pop("size", 10), **kw)  # noqa: E731
thin = Side(style="thin", color="D5DAE0")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# --- Instructions -----------------------------------------------------------
ws = wb.active
ws.title = "Instructions"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 34
ws.column_dimensions["C"].width = 26
ws.column_dimensions["D"].width = 34
ws.column_dimensions["E"].width = 22
row = 1


def put(text, *, bold=False, size=10, color="22282E", col=1):
    global row
    cell = ws.cell(row=row, column=col, value=text)
    cell.font = font(bold=bold, size=size, color=color)
    cell.alignment = Alignment(wrap_text=False, vertical="top")
    row += 1
    return cell


put("ABF-FST Digital Respondent Portal — staff accounts", bold=True, size=14, color=NAVY)
put("Fill in one row per person on the Staff sheet, then send this file to the system administrator.", color="5B6670")
row += 1
put("How to fill it in", bold=True, size=11, color=NAVY)
for line in [
    "1. Go to the Staff sheet. Type only in the yellow cells, one person per row. Don't add notes on that sheet.",
    "2. Full name and Email are required. The email is used for 'Email to me' and to recognise existing accounts.",
    "3. Role: choose from the drop-down list (see the roles below).",
    "4. Provinces: for Contact RAs only. Type one or more provinces separated by commas, or All.",
    "   Where several Contact RAs cover the same province, its cases are shared out evenly between them.",
    "5. Username is optional. Leave it empty and one is made from the name, e.g. tendai.moyo.",
    "6. Do not type passwords. Each person gets a random temporary password, handed to them privately.",
]:
    put(line)
row += 1
put("Example row (copy the format, not the person)", bold=True, size=11, color=NAVY)
for c, value in enumerate(HEADER, start=1):
    cell = ws.cell(row=row, column=c, value=value)
    cell.font = font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.border = box
row += 1
for c, value in enumerate(["Tendai Moyo", "tendai.moyo@example.com", "Contact RA", "Harare, Mashonaland West", ""], start=1):
    cell = ws.cell(row=row, column=c, value=value)
    cell.font = font(italic=True, color="5B6670")
    cell.border = box
row += 2

put("Roles", bold=True, size=11, color=NAVY)
for c, value in enumerate(["Role", "What they do"], start=1):
    cell = ws.cell(row=row, column=c, value=value)
    cell.font = font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
row += 1
for name, what in ROLES:
    ws.cell(row=row, column=1, value=name).font = font(bold=True)
    ws.cell(row=row, column=2, value=what).font = font()
    row += 1
row += 1

put("Main cases per province (to plan Contact RA coverage)", bold=True, size=11, color=NAVY)
for c, value in enumerate(["Province", "Main cases"], start=1):
    cell = ws.cell(row=row, column=c, value=value)
    cell.font = font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
row += 1
first = row
for name, n in CASES:
    ws.cell(row=row, column=1, value=name).font = font()
    ws.cell(row=row, column=2, value=n).font = font(color="0000FF")
    row += 1
ws.cell(row=row, column=1, value="Total").font = font(bold=True)
ws.cell(row=row, column=2, value=f"=SUM(B{first}:B{row - 1})").font = font(bold=True)
row += 1
ws.cell(row=row, column=1, value="Source: live portal register, 15 September 2026. Blue numbers are fixed values, not formulas.").font = font(size=9, italic=True, color="5B6670")

# --- Staff (the sheet that is read) ------------------------------------------
st = wb.create_sheet("Staff")
widths = [28, 34, 30, 40, 22]
for c, (title, width) in enumerate(zip(HEADER, widths), start=1):
    st.column_dimensions[chr(64 + c)].width = width
    cell = st.cell(row=1, column=c, value=title)
    cell.font = font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.border = box
    cell.alignment = Alignment(vertical="center")
st.row_dimensions[1].height = 22
st.freeze_panes = "A2"
ROWS = 80
for r in range(2, ROWS + 2):
    for c in range(1, 6):
        cell = st.cell(row=r, column=c)
        cell.fill = PatternFill("solid", fgColor=INPUT)
        cell.border = box
        cell.font = font()

roles = DataValidation(type="list", formula1='"' + ",".join(name for name, _ in ROLES) + '"', allow_blank=True,
                       showErrorMessage=True, errorTitle="Choose a role", error="Pick a role from the list.")
roles.promptTitle, roles.prompt = "Role", "Choose from the list."
st.add_data_validation(roles)
roles.add(f"C2:C{ROWS + 1}")
provinces = DataValidation(type="custom", formula1="TRUE", allow_blank=True)
provinces.promptTitle = "Provinces (Contact RA only)"
provinces.prompt = "Comma-separated, e.g. Harare, Bulawayo — or All. Leave empty for other roles."
st.add_data_validation(provinces)
provinces.add(f"D2:D{ROWS + 1}")
# No notes below the table: every non-empty row on this sheet is read as a person.

wb.active = 1
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print("wrote", OUT)
