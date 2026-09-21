"""Builds the XLSForm for the KoboToolbox PROIT Interview Profile form from apps/proit/kobo_form.py (the same
definition the portal submits against). Run: python deploy/kobo/build_proit_form.py"""
import importlib.util
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location("kobo_form", ROOT / "backend/apps/proit/kobo_form.py")
F = importlib.util.module_from_spec(spec)
spec.loader.exec_module(F)

OUT = Path(__file__).with_name(f"ABF-FST_PROIT_Interview_Profile_{F.FORM_VERSION_LABEL}_KOBO.xlsx")
wb = openpyxl.Workbook()
survey = wb.active
survey.title = "survey"
survey.append(["type", "name", "label", "required", "appearance", "read_only"])
for meta in ("start", "end", "today", "username"):
    survey.append([meta, meta, "", "", "", ""])
survey.append(["begin_group", F.GROUP_NAME, "Pre-interview profile (sent by the research portal)", "", "field-list", ""])
for name, typ, label, req in F.PROFILE_FIELDS:
    survey.append([typ, name, label, "yes" if req else "", "multiline" if name in (
        "known_evidence_summary", "unresolved_gaps", "contradictions", "priority_probe_questions", "deviation_note") else "", ""])
survey.append(["end_group", "", "", "", "", ""])
survey.append(["begin_repeat", F.REPEAT_NAME, "Facts (one per line of the profile)", "", "", ""])
for name, typ, label, req in F.REPEAT_FIELDS:
    survey.append([typ, name, label, "yes" if req else "", "multiline" if name in ("documentary_value", "sources", "respondent_value", "reconciled_value", "verification_comment") else "", ""])
survey.append(["end_repeat", "", "", "", "", ""])
settings = wb.create_sheet("settings")
settings.append(["form_title", "id_string", "version"])
settings.append([F.FORM_TITLE, "abf_fst_proit_interview_profile", F.FORM_VERSION_LABEL])
readme = wb.create_sheet("README")
for line in [
    "ABF-FST PROIT Interview Profile",
    "Filled in automatically by the research portal when a case's pre-interview profile is reconciled. Do not fill it by hand.",
    "Each record holds one case: what was known before the interview, what the respondent said, and the reconciled value.",
    "Deploy this file as a NEW KoboToolbox project, then set KOBO_PROIT_ASSET_UID on the server (deploy/configure-kobo.sh).",
]:
    readme.append([line])
wb.save(OUT)
print("wrote", OUT)
