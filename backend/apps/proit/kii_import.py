"""Import KII background research done outside the portal (ChatGPT, 2026-10-11).

The PI researches KII organisations in ChatGPT with the "KII research pack" workbook and brings back its three filled
sheets: Fill A Contacts, Fill B PROIT facts, Fill C Interview prep. This loads them into the review flows that already
exist, so nothing ChatGPT found reaches a record until a person accepts it:

- Fill A becomes contact suggestions (ContactSearchRun + ContactProposal), reviewed on the KII record's "Find contact
  details" panel exactly like the portal's own AI contact finder.
- Fill B becomes PROIT proposals (AIResearchRun + AIProposal), reviewed with Accept / Edit / Reject on the record's
  PROIT panel, exactly like the portal's own AI desk research. Only catalogue fields (PROIT_FIELD_CATALOG) are taken,
  and a fact without a source link is refused: accept_proposal() would refuse it anyway.
- Fill C goes into the KII gap-engine notes on the pre-profile (interviewer preparation, never shown to the informant),
  filled where empty and otherwise appended, marked as imported.

A locked pre-profile is never changed. A dry run does everything inside a transaction and rolls it back, so the preview
shows exactly what an import would do.
"""

import re
from io import BytesIO

from django.db import transaction
from django.utils import timezone

from apps.audit.utils import log_action

from .models import (
    PROIT_FIELD_CATALOG,
    AIProposal,
    AIProposalStatus,
    AIResearchRun,
    AIResearchStatus,
    Confidence,
    PreProfile,
    SourceAuthority,
)

MAX_BYTES = 5 * 1024 * 1024
IMPORT_MODEL = "ChatGPT (imported workbook)"
NOT_FOUND = {"", "not found", "n/a", "na", "none", "unknown", "-", "—"}
NOT_PREFILLED = {"organisation_selected", "respondent_selected"}
PROIT_ROLES = {"FINANCE_LENDER", "FARMER_AGRIBUSINESS", "GOVERNMENT_POLICY", "RDC_PROVINCE", "PROCESSOR_AGGREGATOR",
               "INSURANCE_GUARANTEE", "RESEARCH_ACADEMIC"}
URL = re.compile(r"https?://[^\s,;)\]]+")
EMAIL = re.compile(r"^[\w.+-]+@[\w-]+(\.[\w-]+)+$")

# Sheet -> {our key: how the header starts, lower-case}. Matched on the start so small wording changes still work.
SHEETS = {
    "contacts": ("fill a", {
        "kii_id": "kii id", "organisation": "organisation", "website": "website", "main_phone": "main phone",
        "general_email": "general email", "address": "head office address", "person": "person in the role",
        "title": "their title", "person_phone": "their published work phone", "person_email": "their published work email",
        "links": "source links", "confidence": "confidence", "notes": "notes",
    }),
    "facts": ("fill b", {
        "kii_id": "kii id", "organisation": "organisation", "field_id": "field_id", "value": "value found",
        "confidence": "confidence", "title": "source title", "publisher": "publisher", "date": "source date",
        "link": "link", "tier": "source tier", "conflict": "conflicting source", "notes": "notes",
    }),
    "prep": ("fill c", {
        "kii_id": "kii id", "organisation": "organisation", "known": "known evidence", "gaps": "unresolved gaps",
        "contradictions": "contradictions", "probes": "priority probe", "kobo_module": "suggested kii guide module",
        "proit_role": "suggested proit role", "executive": "executive short form", "documents": "public documents",
        "notes": "notes",
    }),
}


class ImportError_(Exception):
    pass


def _text(value) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    return str(value).strip()


def _found(value: str) -> bool:
    return value.strip().lower() not in NOT_FOUND


def read_workbook(data: bytes) -> dict:
    """{"contacts": [...], "facts": [...], "prep": [...]} as dicts keyed as in SHEETS, plus "missing" sheets."""
    from openpyxl import load_workbook

    if len(data) > MAX_BYTES:
        raise ImportError_("The file is larger than 5 MB. Split it, or remove pictures from it.")
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # not an .xlsx, damaged
        raise ImportError_("This isn't an Excel (.xlsx) file the portal can read.") from exc
    out, missing = {}, []
    for key, (prefix, columns) in SHEETS.items():
        sheet = next((ws for ws in wb.worksheets if ws.title.strip().lower().startswith(prefix)), None)
        if sheet is None:
            out[key] = []
            missing.append(prefix.title())
            continue
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            out[key] = []
            continue
        header = [_text(h).lower() for h in rows[0]]
        index = {}
        for ours, starts in columns.items():
            for i, h in enumerate(header):
                if h.startswith(starts) and i not in index.values():
                    index[ours] = i
                    break
        records = []
        for row in rows[1:]:
            item = {ours: _text(row[i]) if i < len(row) else "" for ours, i in index.items()}
            if any(item.values()):
                records.append(item)
        out[key] = records
    if len(missing) == 3:
        raise ImportError_("None of the sheets Fill A Contacts, Fill B PROIT facts or Fill C Interview prep was found.")
    out["missing"] = missing
    return out


def _is_example(item: dict) -> bool:
    return item.get("kii_id", "").upper() in ("", "KII-0000") or "EXAMPLE" in item.get("organisation", "").upper()


def _kii_id(raw: str) -> str:
    digits = re.sub(r"\D", "", raw)
    return f"KII-{int(digits):04d}" if digits else raw.upper()


def _confidence(raw: str) -> str:
    value = raw.strip().upper()
    return value if value in Confidence.values else Confidence.MODERATE


def _tier(raw: str) -> str:
    value = raw.strip().upper()
    if value in SourceAuthority.values:
        return value
    match = re.search(r"TIER[\s_]*([1-4])", value)
    return [v for v in SourceAuthority.values if v.startswith(f"TIER_{match.group(1)}")][0] if match else ""


def _sources(links: str, *, title: str = "", publisher: str = "", published: str = "", authority: str = "") -> list[dict]:
    return [{"title": (title or url)[:400], "url": url[:1000], "publisher": publisher[:255], "published": published,
             "quote": "", "authority": authority} for url in URL.findall(links or "")]


def _profile(record):
    profile = PreProfile.objects.filter(kii_record=record).order_by("id").first()
    if profile is None:
        from .services import create_pre_profile

        profile = create_pre_profile(kii_record=record)
    return profile


def _contacts(record, rows: list[dict], *, user, filename: str) -> tuple[int, list[str]]:
    from apps.contacts.contact_finder import WEBMAIL
    from apps.contacts.models import ContactKind, ContactProposal, ContactSearchRun, ContactSearchStatus

    existing = {(p.kind, p.value.lower(), p.person_name.lower()) for p in record.contact_proposals.all()}
    proposals, warnings = [], []
    for row in rows:
        sources = _sources(row.get("links", ""))
        confidence = _confidence(row.get("confidence", ""))
        before = len(proposals)

        def add(kind, value="", **person):
            key = (kind, value.lower(), person.get("person_name", "").lower())
            if key in existing:
                return
            existing.add(key)
            flags = ["webmail"] if kind == ContactKind.ORG_EMAIL and value.rsplit("@", 1)[-1].lower() in WEBMAIL else []
            proposals.append(dict(kind=kind, value=value[:500], sources=sources, confidence=confidence, flags=flags, **person))

        phone, email = row.get("main_phone", ""), row.get("general_email", "")
        if _found(phone) and len(re.sub(r"\D", "", phone)) >= 9:
            add(ContactKind.ORG_PHONE, phone)
        if _found(email):
            if EMAIL.match(email):
                add(ContactKind.ORG_EMAIL, email)
            else:
                warnings.append(f"general email {email!r} is not an email address")
        if _found(row.get("website", "")):
            add(ContactKind.WEBSITE, row["website"])
        if _found(row.get("address", "")):
            add(ContactKind.OFFICE_LOCATION, row["address"])
        person = row.get("person", "")
        if _found(person):
            p_phone = row.get("person_phone", "") if _found(row.get("person_phone", "")) else ""
            p_email = row.get("person_email", "") if _found(row.get("person_email", "")) and EMAIL.match(row.get("person_email", "")) else ""
            if p_phone or p_email:
                add(ContactKind.PERSON, person_name=person[:255], person_title=row.get("title", "")[:255],
                    person_phone=p_phone[:64], person_email=p_email)
            else:
                # A named person with no published work contact can't be invited; the name still helps PROIT.
                warnings.append(f"{person} has no published work phone or email, so is not a contact suggestion")
        if not sources and len(proposals) > before:
            warnings.append("contacts have no source links: check them before accepting")
    if not proposals:
        return 0, warnings
    run = ContactSearchRun.objects.create(
        kii_record=record, status=ContactSearchStatus.DONE, requested_by=user, provider="import", model=IMPORT_MODEL,
        running_since=timezone.now(), finished_at=timezone.now(),
        summary=f"Imported from {filename} by {getattr(user, 'username', 'unknown')}: found outside the portal (ChatGPT). Check each source before accepting.",
    )
    for p in proposals:
        ContactProposal.objects.create(run=run, kii_record=record, **p)
    return len(proposals), warnings


def _facts(profile, rows: list[dict], *, user, filename: str) -> tuple[int, int, list[str]]:
    grouped: dict[str, list[dict]] = {}
    warnings = []
    for row in rows:
        field_id = row.get("field_id", "").strip()
        if field_id not in PROIT_FIELD_CATALOG or field_id in NOT_PREFILLED:
            warnings.append(f"'{field_id or '(blank)'}' is not a PROIT field the portal pre-fills; row skipped")
            continue
        grouped.setdefault(field_id, []).append(row)
    # Importing the same file twice: a fact already proposed (or decided) with the same value is not proposed again.
    seen = {(p.field_id, p.proposed_value.strip().lower()) for p in profile.ai_proposals.all()}
    grouped = {fid: items for fid, items in grouped.items()
               if not any((fid, r.get("value", "").strip().lower()) in seen for r in items if _found(r.get("value", "")))}
    if not grouped:
        return 0, 0, warnings
    run = AIResearchRun.objects.create(
        pre_profile=profile, status=AIResearchStatus.DONE, requested_by=user, model=IMPORT_MODEL, finished_at=timezone.now(),
        summary=f"Imported from {filename} by {getattr(user, 'username', 'unknown')}: researched outside the portal (ChatGPT). "
                "Open each source before accepting.",
    )
    proposed = not_found = 0
    for field_id, items in grouped.items():
        found = [r for r in items if _found(r.get("value", ""))]
        if not found:
            AIProposal.objects.create(run=run, pre_profile=profile, field_id=field_id, status=AIProposalStatus.NOT_FOUND,
                                      notes="Not found publicly (imported research).")
            not_found += 1
            continue
        sources = []
        for r in found:
            sources += _sources(r.get("link", ""), title=r.get("title", ""), publisher=r.get("publisher", ""),
                                published=r.get("date", ""), authority=_tier(r.get("tier", "")))
        if not sources:
            warnings.append(f"{field_id}: no source link, so it can't be accepted; skipped")
            continue
        values = list(dict.fromkeys(r["value"].strip() for r in found))
        notes = " ".join(r.get("notes", "") for r in found if r.get("notes")).strip()
        if len(values) > 1 or any(r.get("conflict", "").lower().startswith("y") for r in found):
            notes = (f"Sources disagree: {' / '.join(values)}. " + notes).strip()
        AIProposal.objects.create(run=run, pre_profile=profile, field_id=field_id, proposed_value=values[0][:4000],
                                  confidence=_confidence(found[0].get("confidence", "")), sources=sources, notes=notes[:2000])
        proposed += 1
    return proposed, not_found, warnings


def _prep(profile, rows: list[dict], *, filename: str) -> list[str]:
    """Fill the gap-engine notes where empty; append (marked) where something is already written."""
    stamp = f"[Imported from {filename}, {timezone.localdate():%d %b %Y}: check]"
    changed, row = [], {}
    for r in rows:  # several rows for one record: join them
        for key, value in r.items():
            if value:
                row[key] = (row.get(key, "") + "\n" + value).strip() if key in row else value

    def put(attr, text):
        if not _found(text):
            return
        current = getattr(profile, attr)
        setattr(profile, attr, f"{stamp}\n{text}" if not current else f"{current}\n\n{stamp}\n{text}")
        changed.append(attr)

    known = row.get("known", "")
    if _found(row.get("documents", "")):
        known = (known + "\n\nPublic documents to ask about: " + row["documents"]).strip()
    put("known_evidence_summary", known)
    put("unresolved_gaps", row.get("gaps", ""))
    put("contradictions", row.get("contradictions", ""))
    probes = row.get("probes", "")
    if _found(row.get("kobo_module", "")):
        probes = (probes + f"\n\nSuggested KII Guide module(s): {row['kobo_module']}").strip()
    put("priority_probe_questions", probes)
    role = row.get("proit_role", "").strip().upper()
    if role in PROIT_ROLES and not profile.role_specific_module:
        profile.role_specific_module = role
        changed.append("role_specific_module")
    if row.get("executive", "").lower().startswith("y") and not profile.executive_short_form:
        profile.executive_short_form = True
        changed.append("executive_short_form")
    if changed:
        profile.save()
    return changed


def import_workbook(data: bytes, *, filename: str, user, dry_run: bool) -> dict:
    """Read the workbook and load it record by record. With dry_run nothing is kept."""
    from apps.kii.models import KIIRecord

    sheets = read_workbook(data)
    by_record: dict[str, dict] = {}
    skipped_examples = 0
    for key in ("contacts", "facts", "prep"):
        for item in sheets[key]:
            if _is_example(item):
                skipped_examples += 1
                continue
            entry = by_record.setdefault(_kii_id(item["kii_id"]), {"contacts": [], "facts": [], "prep": [], "organisation": ""})
            entry[key].append(item)
            entry["organisation"] = entry["organisation"] or item.get("organisation", "")

    results = []
    with transaction.atomic():
        for kii_id, entry in sorted(by_record.items()):
            result = {"kii_id": kii_id, "record_id": None, "organisation": entry["organisation"], "status": "ready", "contacts": 0,
                      "facts": 0, "not_found": 0, "prep": [], "warnings": []}
            record = KIIRecord.objects.filter(kii_id=kii_id).first()
            if record is None:
                result.update(status="skipped", warnings=[f"No KII record {kii_id} in the portal"])
                results.append(result)
                continue
            from apps.contacts.kii_finder import organisation_name

            result["organisation"], result["record_id"] = organisation_name(record) or entry["organisation"], record.id
            result["contacts"], warnings = _contacts(record, entry["contacts"], user=user, filename=filename)
            result["warnings"] += warnings
            if entry["facts"] or entry["prep"]:
                profile = _profile(record)
                if profile.prepopulation_locked_at is not None:
                    result["warnings"].append("PROIT profile is locked, so facts and interview notes were not imported")
                else:
                    result["facts"], result["not_found"], warnings = _facts(profile, entry["facts"], user=user, filename=filename)
                    result["warnings"] += warnings
                    result["prep"] = _prep(profile, entry["prep"], filename=filename)
            if not (result["contacts"] or result["facts"] or result["not_found"] or result["prep"]):
                result["status"] = "nothing new"
            elif not dry_run:
                result["status"] = "imported"
                log_action("proit.kii_research_imported", record, {
                    "kii_id": kii_id, "file": filename[:120], "contacts": result["contacts"], "facts": result["facts"],
                    "not_found": result["not_found"], "prep": result["prep"], "user_id": getattr(user, "id", None),
                }, user=user)
            results.append(result)
        if dry_run:
            transaction.set_rollback(True)

    totals = {key: sum(r[key] for r in results) for key in ("contacts", "facts", "not_found")}
    totals["records"] = sum(1 for r in results if r["status"] in ("ready", "imported"))
    return {"dry_run": dry_run, "file": filename, "records": results, "totals": totals,
            "examples_skipped": skipped_examples, "missing_sheets": sheets["missing"]}
