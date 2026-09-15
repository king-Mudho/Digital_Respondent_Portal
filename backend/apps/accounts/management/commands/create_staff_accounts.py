"""
Create named staff accounts from the PI's spreadsheet, and optionally split
the Main cases between the new Contact RAs (added 2026-09-15).

Until now every role used one shared account (contact_ra, kii_ra, ...), so
nothing in the audit log could say which person did what, and all 400 Main
cases sat on a single Contact RA.

The spreadsheet is docs/templates/ABF-FST_Staff_Accounts_Template.xlsx: one
person per row -- full name, email, role, and for Contact RAs the provinces
they cover ("All" for everywhere). Each new account gets a random temporary
password, written once to a private credentials file for the PI to hand out;
it is never printed or logged. Existing accounts (same email or username)
are left untouched.

With --assign-cases, Main cases still on the shared account (or unassigned)
are divided evenly, by Sample ID, among the Contact RAs covering each
province. Cases already given to a named RA never move.

    manage.py create_staff_accounts --file staff.xlsx --dry-run
    manage.py create_staff_accounts --file staff.xlsx --assign-cases --credentials-dir /root
"""

import csv
import re
import secrets
import unicodedata
from collections import defaultdict
from pathlib import Path

from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.core.validators import validate_email
from django.db import transaction
from django.utils import timezone

from apps.accounts.models import Role, User
from apps.audit.utils import log_action
from apps.sampling.models import Province, SampleCase, SampleType
from apps.sampling.services import bulk_assign_cases

ROLE_ALIASES = {
    **{code.lower(): code for code, _ in Role.NAME_CHOICES},
    **{label.lower(): code for code, label in Role.NAME_CHOICES},
    "pi": Role.PI_ADMIN, "pi / admin": Role.PI_ADMIN, "admin": Role.PI_ADMIN,
    "field coordinator": Role.FIELD_COORDINATOR, "digital coordinator": Role.FIELD_COORDINATOR,
    "coordinator": Role.FIELD_COORDINATOR,
    "qa ra": Role.QUAN_QA_RA, "quan qa ra": Role.QUAN_QA_RA, "questionnaire qa ra": Role.QUAN_QA_RA,
    "documents ra": Role.DOCUMENTARY_RA, "document ra": Role.DOCUMENTARY_RA,
    "analyst": Role.ANALYST, "supervisor": Role.SUPERVISOR_READONLY,
}
PROVINCES = {**{label.lower(): code for code, label in Province.choices}, **{code.lower(): code for code, _ in Province.choices}}
SIGN_IN = "https://research.agribizframework.com/admin/login"


class _Run:
    pk = "create_staff_accounts"


def _cell(value) -> str:
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def read_rows(path: Path) -> list[dict]:
    """Rows as {full_name, email, role, provinces, username} from .xlsx or .csv."""
    if path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8-sig") as f:
            table = [list(r) for r in csv.reader(f)]
    else:
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)
        ws = wb["Staff"] if "Staff" in wb.sheetnames else wb.worksheets[0]
        table = [list(r) for r in ws.iter_rows(values_only=True)]
    header_at = next((i for i, r in enumerate(table) if any(_cell(v).lower() == "full name" for v in r)), None)
    if header_at is None:
        raise CommandError("No header row with a 'Full name' column was found.")
    header = [_cell(v).lower() for v in table[header_at]]

    def col(*names):
        return next((i for i, h in enumerate(header) if any(h.startswith(n) for n in names)), None)

    idx = {"full_name": col("full name"), "email": col("email"), "role": col("role"),
           "provinces": col("provinces"), "username": col("username")}
    if None in (idx["full_name"], idx["email"], idx["role"]):
        raise CommandError("The sheet needs Full name, Email and Role columns.")
    rows = []
    for number, r in enumerate(table[header_at + 1:], start=header_at + 2):
        values = {k: (_cell(r[i]) if i is not None and i < len(r) else "") for k, i in idx.items()}
        if any(values.values()):
            rows.append({**values, "row": number})
    return rows


def make_username(full_name: str, taken: set[str]) -> str:
    ascii_name = unicodedata.normalize("NFKD", full_name).encode("ascii", "ignore").decode()
    parts = [p for p in re.split(r"[^a-z]+", ascii_name.lower()) if p and p not in {"mr", "mrs", "ms", "dr", "prof"}]
    base = ".".join([parts[0], parts[-1]] if len(parts) > 1 else parts) or "staff"
    username, n = base, 2
    while username in taken:
        username, n = f"{base}{n}", n + 1
    return username


def temporary_password(user: User) -> str:
    while True:
        password = "-".join(secrets.token_urlsafe(4) for _ in range(3))
        try:
            validate_password(password, user=user)
            return password
        except ValidationError:
            continue


class Command(BaseCommand):
    help = "Create named staff accounts from a spreadsheet and optionally split Main cases between Contact RAs."

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True, help=".xlsx (Staff sheet) or .csv")
        parser.add_argument("--dry-run", action="store_true", help="Check the sheet and show the plan; change nothing.")
        parser.add_argument("--assign-cases", action="store_true", help="Split Main cases among the Contact RAs by province.")
        parser.add_argument("--from-user", default="contact_ra", help="Shared account whose cases are split (default contact_ra).")
        parser.add_argument("--credentials-dir", default=".", help="Where to write the one-time credentials file.")

    def handle(self, *args, file, dry_run=False, assign_cases=False, from_user="contact_ra", credentials_dir=".", **options):
        path = Path(file)
        if not path.exists():
            raise CommandError(f"File not found: {path}")
        rows = read_rows(path)
        if not rows:
            raise CommandError("The sheet has no people in it.")

        problems, plan, seen_emails = [], [], set()
        taken = set(User.objects.values_list("username", flat=True))
        for r in rows:
            where = f"row {r['row']}"
            role = ROLE_ALIASES.get(r["role"].lower())
            if not r["full_name"]:
                problems.append(f"{where}: Full name is empty.")
            try:
                validate_email(r["email"])
            except ValidationError:
                problems.append(f"{where}: '{r['email']}' is not a valid email address.")
            if role is None:
                problems.append(f"{where}: unknown role '{r['role']}'.")
            if r["email"].lower() in seen_emails:
                problems.append(f"{where}: {r['email']} appears twice in the sheet.")
            seen_emails.add(r["email"].lower())

            provinces = []
            for name in filter(None, (p.strip() for p in re.split(r"[,;/]", r["provinces"]))):
                if name.lower() in ("all", "all provinces"):
                    provinces = [code for code, _ in Province.choices]
                elif name.lower() in PROVINCES:
                    provinces.append(PROVINCES[name.lower()])
                else:
                    problems.append(f"{where}: unknown province '{name}'.")
            if provinces and role != Role.CONTACT_RA:
                problems.append(f"{where}: provinces are only for Contact RAs; leave the column empty for {r['role']}.")

            existing = User.objects.filter(email__iexact=r["email"]).first() or (
                User.objects.filter(username=r["username"]).first() if r["username"] else None)
            username = existing.username if existing else (r["username"] or make_username(r["full_name"], taken))
            taken.add(username)
            plan.append({**r, "role_code": role, "province_codes": sorted(set(provinces)), "username": username, "existing": existing})

        if problems:
            raise CommandError("Fix these in the sheet, then run again:\n  " + "\n  ".join(problems))

        for p in plan:
            state = f"exists already ({p['existing'].username}), left unchanged" if p["existing"] else "new"
            cover = f" · provinces: {', '.join(p['province_codes'])}" if p["province_codes"] else ""
            self.stdout.write(f"{p['full_name']} <{p['email']}> -> {p['username']} [{p['role_code']}] {state}{cover}")

        split = self._case_split(plan, from_user) if assign_cases else {}
        for username, ids in split.items():
            self.stdout.write(f"  cases for {username}: {len(ids)}")
        if assign_cases:
            covered = {code for p in plan for code in p["province_codes"]}
            uncovered = SampleCase.objects.filter(sample_type=SampleType.MAIN).exclude(organisation__province__in=covered).count()
            if uncovered:
                self.stdout.write(self.style.WARNING(f"  {uncovered} Main cases are in provinces no Contact RA covers; they stay where they are."))

        new = [p for p in plan if not p["existing"]]
        self.stdout.write(self.style.SUCCESS(f"{len(new)} to create, {len(plan) - len(new)} already exist."
                                             + (" (DRY RUN -- nothing changed)" if dry_run else "")))
        if dry_run:
            return

        credentials = []
        with transaction.atomic():
            roles = {r.name: r for r in Role.objects.all()}
            for p in new:
                first, _, last = p["full_name"].partition(" ")
                user = User(username=p["username"], email=p["email"], first_name=first, last_name=last,
                            role=roles[p["role_code"]], is_staff=p["role_code"] == Role.PI_ADMIN, is_active=True)
                password = temporary_password(user)
                user.set_password(password)
                user.save()
                credentials.append([p["full_name"], p["email"], dict(Role.NAME_CHOICES)[p["role_code"]], user.username, password, SIGN_IN])
            for username, ids in split.items():
                bulk_assign_cases(to_user=User.objects.get(username=username), sample_ids=ids)
            log_action("accounts.staff_accounts_created", _Run(), {
                "created": [{"username": p["username"], "role": p["role_code"]} for p in new],
                "already_existed": [p["username"] for p in plan if p["existing"]],
                "cases_assigned": {u: len(ids) for u, ids in split.items()},
            })

        if credentials:
            out = Path(credentials_dir) / f"staff-credentials-{timezone.localtime():%Y%m%d-%H%M%S}.csv"
            with out.open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(["Full name", "Email", "Role", "Username", "Temporary password", "Sign in at"])
                writer.writerows(credentials)
            out.chmod(0o600)
            self.stdout.write(self.style.SUCCESS(
                f"Created {len(credentials)} accounts. Temporary passwords are in {out} -- give each person theirs "
                "privately, ask them to change it at first sign-in, then delete the file."))
        else:
            self.stdout.write(self.style.SUCCESS("No new accounts were needed."))

    def _case_split(self, plan, from_username) -> dict[str, list[str]]:
        """Sample IDs per Contact RA: each province's movable cases dealt out
        in turn among the RAs covering it."""
        shared = User.objects.filter(username=from_username).first()
        covering = defaultdict(list)
        for p in plan:
            if p["role_code"] == Role.CONTACT_RA:
                for code in p["province_codes"]:
                    covering[code].append(p["username"])
        split = defaultdict(list)
        for province, usernames in covering.items():
            movable = SampleCase.objects.filter(sample_type=SampleType.MAIN, organisation__province=province)
            movable = movable.filter(assigned_ra__isnull=True) | (movable.filter(assigned_ra=shared) if shared else movable.none())
            for i, sample_id in enumerate(movable.order_by("sample_id").values_list("sample_id", flat=True)):
                split[usernames[i % len(usernames)]].append(sample_id)
        return dict(split)
