"""
Full KoboToolbox data for the PI: an analysis-ready workbook per form, and
every completed form as a PDF in one ZIP (added 2026-09-15).

The portal's other exports hold portal metadata only (case, QA state,
timings) -- never the answers. Cleaning and analysis need the answers, with
both the stored codes (for SPSS/Stata/R) and the readable labels, plus a data
dictionary, so that is what this builds, straight from the deployed form and
the live KoboToolbox data.

PI only: raw answers can identify organisations and, in the KII Guide, people.
Every export is audited.
"""

import csv
import io
import zipfile
from datetime import datetime

from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from apps.audit.utils import log_action

from . import submission_copies as copies
from .client import KoboClient
from .pdf import _Form, _label, _type_parts, render_submission_pdf

META_COLUMNS = ["_id", "_uuid", "meta/rootUuid", "_submission_time", "_submitted_by", "start", "end", "today"]
NON_ANSWER_TYPES = {"note", "begin_group", "end_group", "begin_repeat", "end_repeat", "start", "end", "today",
                    "username", "deviceid", "audit"}
NUMERIC_TYPES = {"integer", "decimal", "range"}
PORTAL_COLUMNS = ["portal_matched_case", "portal_qa_status", "portal_consent_withdrawn", "portal_case_status"]


class _Ref:
    def __init__(self, key):
        self.pk = key


def _survey_columns(content: dict):
    """(top-level columns, {repeat xpath: (label, [columns])}) in form order.
    A column is (xpath, row)."""
    top, repeats, stack = [], {}, []
    for row in content.get("survey", []):
        kind = (row.get("type") or "").strip().replace(" ", "_")
        base, _ = _type_parts(row)
        if kind == "begin_repeat":
            stack.append(row.get("$xpath") or row.get("name"))
            repeats[stack[-1]] = (_label(row), [])
            continue
        if kind == "end_repeat":
            stack.pop() if stack else None
            continue
        if kind in NON_ANSWER_TYPES or base in NON_ANSWER_TYPES:
            continue
        xpath = row.get("$xpath") or row.get("name")
        (repeats[stack[-1]][1] if stack else top).append((xpath, row))
    return top, repeats


def _code(row: dict, value):
    if value in (None, ""):
        return None
    base, _ = _type_parts(row)
    text = str(value)
    if base in NUMERIC_TYPES or (base == "select_one" and text.lstrip("-").isdigit()):
        try:
            return int(text)
        except ValueError:
            try:
                return float(text)
            except ValueError:
                return text
    return text


def _cell(ws, value, bold=False):
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"  # never let an answer run as a spreadsheet formula
    if bold:
        cell.font = Font(bold=True)
    return cell


def _portal_columns(payloads: list[dict]) -> dict[int, list]:
    """Questionnaire only: what the portal knows about each submission's case."""
    from apps.consent.models import ConsentDecision, ConsentType
    from apps.consent.services import latest_consent
    from apps.sampling.models import SampleCase

    from .models import QUANSubmission
    from .services import _payload_sample_id, _submission_identity

    identities = {_submission_identity(p): p.get("_id") for p in payloads}
    by_identity = {s.kobo_submission_uuid: s for s in QUANSubmission.objects.filter(
        kobo_submission_uuid__in=[i for i in identities if i]).select_related("sample_case")}
    sample_ids = {_payload_sample_id(p) for p in payloads}
    cases = {c.sample_id: c for c in SampleCase.objects.filter(sample_id__in=[s for s in sample_ids if s])}
    out = {}
    for payload in payloads:
        submission = by_identity.get(_submission_identity(payload))
        case = submission.sample_case if submission else cases.get(_payload_sample_id(payload))
        consent = latest_consent(case, ConsentType.PARTICIPATION) if case else None
        out[payload.get("_id")] = [
            case.sample_id if submission else "",
            submission.qa_status if submission else "NOT_IN_PORTAL",
            (consent is not None and consent.decision == ConsentDecision.WITHDRAWN) if case else "",
            case.workflow_status if case else "",
        ]
    return out


def build_workbook(key: str, *, user) -> tuple[bytes, str]:
    copies.require_form(key, user)
    content = copies.form_content(key)
    payloads = KoboClient(asset_uid=copies.asset_uid(key)).fetch_submissions()
    form = _Form(content)
    top, repeats = _survey_columns(content)
    portal = _portal_columns(payloads) if key == "questionnaire" else {}
    withdrawn = sum(1 for values in portal.values() if values[2] is True)

    wb = Workbook(write_only=True)
    readme = wb.create_sheet("README")
    exported = timezone.localtime()
    for line in [
        [copies.FORMS[key]["title"]],
        [f"Exported {exported:%d %b %Y %H:%M} (Africa/Harare) from KoboToolbox via the ABF-FST portal"],
        [f"{len(payloads)} submissions"],
        [],
        ["Sheet", "What it holds"],
        ["data_codes", "One row per submission. Answers as stored codes -- use for SPSS, Stata or R."],
        ["data_labels", "The same rows with choice codes replaced by their labels."],
        ["questions", "Data dictionary: column name, question text, type, choice list, required."],
        ["choices", "Every choice list: code and label."],
        *[[f"repeat_{i + 1}", f"{label}: one row per entry, linked by _parent_id"] for i, (label, _) in enumerate(repeats.values())],
        [],
        ["Columns starting with _ are KoboToolbox metadata. meta/rootUuid is the submission's stable id across edits."],
        *([["portal_* columns: the matched case, its QA state, whether consent was later withdrawn, and its workflow status."],
           [f"WITHDRAWN PARTICIPANTS: {withdrawn}. Leave every row with portal_consent_withdrawn = TRUE out of analysis "
            "(PI decision, 15 Sep 2026: answers kept for the audit trail, never analysed)."]]
          if key == "questionnaire" else []),
        ["Confidential research data. Store and share only as the approved data management plan allows."],
    ]:
        readme.append([_cell(readme, v, bold=(line and line[0] in ("Sheet",))) for v in line])

    headers = META_COLUMNS + [xpath for xpath, _ in top] + (PORTAL_COLUMNS if portal else [])
    for sheet_name, as_labels in (("data_codes", False), ("data_labels", True)):
        ws = wb.create_sheet(sheet_name)
        ws.append([_cell(ws, h, bold=True) for h in headers])
        for payload in payloads:
            row = [payload.get(c) for c in META_COLUMNS]
            for xpath, q in top:
                value = payload.get(xpath)
                row.append(None if value in (None, "") else (form.answer_text(q, value) if as_labels else _code(q, value)))
            row += portal.get(payload.get("_id"), [])
            ws.append([_cell(ws, v) for v in row])

    for i, (repeat_xpath, (label, columns)) in enumerate(repeats.items()):
        ws = wb.create_sheet(f"repeat_{i + 1}")
        ws.append([_cell(ws, h, bold=True) for h in ["_parent_id", "_index"] + [x for x, _ in columns]])
        for payload in payloads:
            for n, entry in enumerate(payload.get(repeat_xpath) or [], start=1):
                ws.append([_cell(ws, v) for v in
                           [payload.get("_id"), n] + [_code(q, entry.get(x)) for x, q in columns]])

    questions = wb.create_sheet("questions")
    questions.append([_cell(questions, h, bold=True) for h in ["column", "question", "type", "choice_list", "required", "repeat"]])
    for xpath, q in top + [c for _, cols in repeats.values() for c in cols]:
        base, list_name = _type_parts(q)
        questions.append([_cell(questions, v) for v in [xpath, _label(q), base, list_name or "", bool(q.get("required")),
                                                           next((lbl for rx, (lbl, cols) in repeats.items() if (xpath, q) in cols), "")]])

    choices = wb.create_sheet("choices")
    choices.append([_cell(choices, h, bold=True) for h in ["choice_list", "code", "label"]])
    for c in content.get("choices", []):
        choices.append([_cell(choices, v) for v in [c.get("list_name"), c.get("name") or c.get("$autovalue"), _label(c)]])

    buffer = io.BytesIO()
    wb.save(buffer)
    log_action("kobo.data_exported", _Ref(key), {"form": key, "format": "xlsx", "submissions": len(payloads),
                                                 "user_id": user.id}, user=user)
    return buffer.getvalue(), f"{slugify(key)}-data-{exported:%Y%m%d-%H%M}.xlsx"


class _Sink:
    """A write-only file for zipfile that hands back what was written, so the
    ZIP streams out as each PDF is added instead of being built in memory
    first (and outliving a proxy timeout)."""

    def __init__(self):
        self.chunks: list[bytes] = []

    def write(self, data):
        self.chunks.append(bytes(data))
        return len(data)

    def flush(self):
        pass

    def drain(self) -> bytes:
        data, self.chunks = b"".join(self.chunks), []
        return data


def stream_pdf_zip(key: str, *, user):
    """(generator of bytes, filename). Checks access before streaming starts."""
    copies.require_form(key, user)
    content = copies.form_content(key)
    payloads = KoboClient(asset_uid=copies.asset_uid(key)).fetch_submissions()
    title = copies.FORMS[key]["title"]
    stamp = datetime.now().strftime("%Y%m%d-%H%M")
    # Withdrawn participants' forms stay in the archive, visibly marked (PI decision, 15 Sep 2026).
    portal = _portal_columns(payloads) if key == "questionnaire" else {}

    def generate():
        sink = _Sink()
        manifest = io.StringIO()
        writer = csv.writer(manifest)
        writer.writerow(["file", "record", "kobo_id", "submitted_utc", "consent_withdrawn"])
        with zipfile.ZipFile(sink, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            seen = set()
            for payload in payloads:
                label = copies.record_label(key, payload)
                withdrawn = portal.get(payload.get("_id"), [None, None, None])[2] is True
                name = f"{'WITHDRAWN-' if withdrawn else ''}{slugify(label) or 'record'}-{payload.get('_id')}.pdf"
                if name in seen:
                    continue
                seen.add(name)
                pdf = render_submission_pdf(form_content=content, payload=payload, form_title=title, record_label=label)
                archive.writestr(name, pdf)
                writer.writerow([name, label, payload.get("_id"), payload.get("_submission_time"), withdrawn])
                yield sink.drain()
            archive.writestr("manifest.csv", manifest.getvalue())
        yield sink.drain()
        log_action("kobo.data_exported", _Ref(key), {"form": key, "format": "pdf_zip", "submissions": len(payloads),
                                                     "user_id": user.id}, user=user)

    return generate(), f"{slugify(key)}-completed-forms-{stamp}.zip"
