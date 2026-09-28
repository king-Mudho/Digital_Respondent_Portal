"""
Analysis packs: the portal's data, ready to open in the two tools the study analyses it with.

  smartpls    the questionnaire for SmartPLS 4 -- one row per accepted respondent, numeric only, short variable
              names, a missing-value code, a data dictionary. No names, ids, contact details or free text.
  atlas-kii   the key informant interviews for ATLAS.ti -- one Word document per interview (question, then answer)
              plus an Excel sheet of interview attributes for document groups. No organisation, job title or contact.
  atlas-documents
              the uploaded source documents for ATLAS.ti, named by DOC-ID, plus an attribute sheet.

Built from the deployed forms themselves (labels, choice lists), so a revised form needs no code change, and from the
live KoboToolbox data -- the same source as data_export.py. PI only, every download audited: raw answers can identify
organisations and, in the KII Guide, people.

Rules carried over from the rest of the portal: a participant who withdrew consent is left out (PI decision, 15 Sep
2026), and only questionnaire submissions that passed QA are analysed unless the caller asks for the rest.
"""

import csv
import io
import re
import tempfile
import zipfile
from xml.sax.saxutils import escape

import requests
from django.utils import timezone
from django.utils.text import slugify
from openpyxl import Workbook

from apps.audit.utils import log_action

from . import submission_copies as copies
from .client import KoboClient
from .data_export import _portal_columns
from .pdf import _Form, _label, _type_parts

MISSING = -99  # the missing-value marker written into the SmartPLS file; the README tells the analyst to enter it
# Choices that mean "no answer" rather than a category: coded as missing, and left out of the 1..n numbering.
MISSING_CHOICES = {"na", "not_sure_na", "prefer_not"}
NUMERIC_BASES = {"integer", "decimal", "range"}
CHOICE_BASES = {"select_one", "select_multiple"}

# KII fields kept out of the interview documents and the attribute sheet: they name or locate the person, or are
# administrative. (The PI's full Kobo workbook still has them.)
KII_LEFT_OUT = {
    "KII_ID", "org_represented", "official_title", "consent", "recording_consent", "start_time", "end_time",
    "may_contact", "preferred_channel", "followup_referral", "interviewer_code",
}
# Structured facts about each interview, for ATLAS.ti document groups.
KII_ATTRIBUTES = ["interview_date", "RESP_CAT", "RESP_CAT_OTHER", "MODE", "FORMAT", "CAPTURE_MODE", "duration_minutes",
                  "COMPLETION_STATUS", "ROLE_MODULES", "documents_available"]


class _Ref:
    def __init__(self, key):
        self.pk = key


def _fetch(key: str, user):
    """(form content, submissions) for one Kobo form, with a plain error if KoboToolbox can't be reached."""
    copies.require_form(key, user)
    try:
        content = copies.form_content(key)
        payloads = KoboClient(asset_uid=copies.asset_uid(key)).fetch_submissions()
    except requests.RequestException as exc:
        raise copies.CopyError("kobo_unreachable", f"KoboToolbox couldn't be reached ({exc.__class__.__name__}).", 502) from exc
    return content, payloads


def _walk(content: dict):
    """(section label, xpath, row) for every answer question outside a repeat group, in form order."""
    groups, repeat_depth = [], 0
    for row in content.get("survey", []):
        kind = (row.get("type") or "").strip().replace(" ", "_")
        if kind == "begin_group":
            groups.append(_label(row))
        elif kind == "end_group":
            groups and groups.pop()
        elif kind == "begin_repeat":
            repeat_depth += 1
        elif kind == "end_repeat":
            repeat_depth -= 1
        elif repeat_depth == 0 and _type_parts(row)[0] in (CHOICE_BASES | NUMERIC_BASES | {"text", "date", "time", "datetime"}):
            yield (groups[0] if groups else ""), row.get("$xpath") or row.get("name"), row


def _choice_lists(content: dict) -> dict[str, list[str]]:
    lists: dict[str, list[str]] = {}
    for c in content.get("choices", []):
        lists.setdefault(c.get("list_name"), []).append(str(c.get("name") or c.get("$autovalue")))
    return lists


def _coding(names: list[str]) -> dict[str, int]:
    """Numeric code for each choice of one list: numbers stay as they are, yes/no become 1/0, anything else is
    numbered in the order the form lists it, with 'no answer' choices set aside as missing."""
    if names and all(re.fullmatch(r"-?\d+", n) for n in names):
        return {n: int(n) for n in names}
    if set(names) == {"yes", "no"}:
        return {"yes": 1, "no": 0}
    coded, k = {}, 0
    for n in names:
        if n in MISSING_CHOICES:
            coded[n] = MISSING
        else:
            k += 1
            coded[n] = k
    return coded


def _var(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9_]", "_", str(name))


def _number(value):
    try:
        f = float(str(value))
    except (TypeError, ValueError):
        return MISSING
    return int(f) if f.is_integer() else f


def _zip_response_file(entries: dict[str, bytes | str]):
    """A temp file holding the ZIP (kept off memory: the documents pack can be large)."""
    tmp = tempfile.TemporaryFile()
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries.items():
            z.writestr(name, data)
    tmp.seek(0)
    return tmp


def _csv_bytes(header: list, rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8")


def _xlsx_bytes(sheets: dict[str, tuple[list, list[list]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, (header, rows) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for r in rows:
            ws.append(r)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str) and cell.value[:1] in ("=", "+", "-", "@"):
                    cell.data_type = "s"  # never let an answer run as a spreadsheet formula
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _stamp() -> str:
    return timezone.localtime().strftime("%Y%m%d-%H%M")


# ---------------------------------------------------------------------------------------------- SmartPLS 4

def build_smartpls_pack(*, user, include_unreviewed: bool = False):
    """(temp ZIP file, filename, summary) -- see the module docstring."""
    content, payloads = _fetch("questionnaire", user)
    portal = _portal_columns(payloads)
    lists = _choice_lists(content)
    form = _Form(content)

    kept, withdrawn, not_passed = [], 0, 0
    for p in payloads:
        sample_id, qa_status, was_withdrawn, _ = portal.get(p.get("_id"), ["", "NOT_IN_PORTAL", "", ""])
        if was_withdrawn is True:
            withdrawn += 1
        elif qa_status != "QA_PASSED" and not include_unreviewed:
            not_passed += 1
        else:
            kept.append((p, sample_id))
    kept.sort(key=lambda t: (t[0].get("_submission_time") or "", t[0].get("_id") or 0))

    columns, dictionary, left_out = [], [], []   # columns: (variable, reader)
    for section, xpath, row in _walk(content):
        base, list_name = _type_parts(row)
        name, question = row.get("name"), _label(row)
        if base in NUMERIC_BASES:
            columns.append((_var(name), (lambda p, x=xpath: _number(p.get(x)) if p.get(x) not in (None, "") else MISSING)))
            dictionary.append([_var(name), question, base, "numeric answer as entered", f"blank = {MISSING}", section])
        elif base == "select_one":
            names = lists.get(list_name, [])
            coding = _coding(names)
            columns.append((_var(name), (lambda p, x=xpath, c=coding: c.get(str(p.get(x)), MISSING) if p.get(x) not in (None, "") else MISSING)))
            shown = "; ".join(f"{v}={form.choices.get((list_name, k), k)}" for k, v in coding.items() if v != MISSING)
            dictionary.append([_var(name), question, "select_one", shown, f"blank, not applicable or prefer not to say = {MISSING}", section])
        elif base == "select_multiple":
            for choice in lists.get(list_name, []):
                var = _var(f"{name}_{choice}")
                columns.append((var, (lambda p, x=xpath, ch=choice: (MISSING if p.get(x) in (None, "") else int(ch in str(p.get(x)).split())))))
                dictionary.append([var, f"{question} [{form.choices.get((list_name, choice), choice)}]", "select_multiple (one column per choice)",
                                   "1 = chosen, 0 = not chosen", f"question not answered = {MISSING}", section])
        else:
            left_out.append([name, question, base])

    header = ["respondent_no"] + [v for v, _ in columns]
    rows = [[i] + [read(p) for _, read in columns] for i, (p, _) in enumerate(kept, start=1)]
    key_rows = [[i, sample_id, p.get("_id"), p.get("_submission_time")] for i, (p, sample_id) in enumerate(kept, start=1)]

    readme = "\r\n".join([
        "ABF-FST questionnaire, prepared for SmartPLS 4",
        f"Exported {timezone.localtime():%d %b %Y %H:%M} (Africa/Harare) from KoboToolbox through the ABF-FST portal.",
        "",
        f"Respondents in the file: {len(rows)}",
        f"Left out: {withdrawn} who later withdrew consent (kept in the audit trail, never analysed)"
        + ("" if include_unreviewed else f"; {not_passed} not yet accepted by QA."),
        "",
        "FILES",
        "  data.csv / data.xlsx      the same data. One row per respondent, one column per variable, first row = variable names.",
        "  data_dictionary.csv       what every variable is, how it is coded, and what is missing.",
        "  PRIVATE_do_not_share/     respondent_key.csv links respondent_no back to the Sample ID. Never share it or import it.",
        "",
        "IMPORTING IN SMARTPLS 4",
        "  1. Use New project, then Import data file, and choose data.csv (comma separated) or data.xlsx.",
        f"  2. Enter {MISSING} as the missing-value marker when asked, or set it in the data view afterwards.",
        "  3. Every variable is numeric. Use the data dictionary to pick indicators and to find the codes.",
        "  4. respondent_no is only a row number: do not use it as an indicator.",
        "",
        "HOW ANSWERS ARE CODED",
        "  Questions stored as numbers in the form (the 1-5 scales) keep those numbers.",
        "  Yes/no become 1/0. Other single-choice questions are numbered 1, 2, 3... in the order the form lists them",
        "  (see the dictionary). 'Not applicable', 'not sure' and 'prefer not to say' are missing, not a category.",
        "  Multiple-choice questions become one 0/1 column per choice.",
        "  Free-text answers, dates, ids and contact details are not in this file. Text goes to ATLAS.ti.",
        "",
        "Confidential research data. Store and share only as the approved data management plan allows.",
    ])
    entries = {
        "data.csv": _csv_bytes(header, rows),
        "data.xlsx": _xlsx_bytes({"data": (header, rows),
                                  "dictionary": (["variable", "question", "type", "coding", "missing", "section"], dictionary)}),
        "data_dictionary.csv": _csv_bytes(["variable", "question", "type", "coding", "missing", "section"], dictionary),
        "README.txt": readme,
        "PRIVATE_do_not_share/respondent_key.csv": _csv_bytes(["respondent_no", "sample_id", "kobo_id", "submitted"], key_rows),
    }
    summary = {"pack": "smartpls", "respondents": len(rows), "variables": len(columns), "withdrawn_left_out": withdrawn,
               "not_passed_left_out": 0 if include_unreviewed else not_passed, "text_or_other_left_out": len(left_out)}
    return _zip_response_file(entries), f"ABF-FST_SmartPLS_{_stamp()}.zip", summary


# ------------------------------------------------------------------------------------------- ATLAS.ti: KII

_DOCX_TYPES = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    '<Override PartName="/word/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/></Types>')
_DOCX_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
_DOCX_DOC_RELS = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                  '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
_W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
_DOCX_STYLES = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:styles {_W}>'
    '<w:docDefaults><w:rPrDefault><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/><w:sz w:val="22"/></w:rPr></w:rPrDefault>'
    '<w:pPrDefault><w:pPr><w:spacing w:after="120" w:line="276" w:lineRule="auto"/></w:pPr></w:pPrDefault></w:docDefaults>'
    '<w:style w:type="paragraph" w:default="1" w:styleId="Normal"><w:name w:val="Normal"/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Title"><w:name w:val="Title"/><w:basedOn w:val="Normal"/><w:pPr><w:spacing w:after="200"/></w:pPr><w:rPr><w:b/><w:sz w:val="32"/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading1"><w:name w:val="heading 1"/><w:basedOn w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="280" w:after="100"/><w:outlineLvl w:val="0"/></w:pPr><w:rPr><w:b/><w:sz w:val="26"/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="Heading2"><w:name w:val="heading 2"/><w:basedOn w:val="Normal"/><w:pPr><w:keepNext/><w:spacing w:before="200" w:after="60"/><w:outlineLvl w:val="1"/></w:pPr><w:rPr><w:b/><w:sz w:val="22"/></w:rPr></w:style>'
    '</w:styles>')
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _para(style: str | None, text: str) -> str:
    ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
    return f'<w:p>{ppr}<w:r><w:t xml:space="preserve">{escape(_CONTROL.sub("", text))}</w:t></w:r></w:p>'


def docx_bytes(blocks: list[tuple[str | None, str]]) -> bytes:
    """A small, valid .docx from (style, text) blocks; style is Title, Heading1, Heading2 or None for body text."""
    body = "".join(_para(style, line) for style, text in blocks for line in (str(text).splitlines() or [""]))
    document = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {_W}><w:body>{body}</w:body></w:document>'
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", _DOCX_TYPES)
        z.writestr("_rels/.rels", _DOCX_RELS)
        z.writestr("word/document.xml", document)
        z.writestr("word/styles.xml", _DOCX_STYLES)
        z.writestr("word/_rels/document.xml.rels", _DOCX_DOC_RELS)
    return out.getvalue()


def _kii_withdrawn(kii_id: str) -> bool:
    from apps.consent.models import ConsentDecision, ConsentType
    from apps.consent.services import latest_consent
    from apps.kii.models import KIIRecord

    record = KIIRecord.objects.filter(kii_id=kii_id).first()
    consent = latest_consent(record, ConsentType.PARTICIPATION) if record else None
    return consent is not None and consent.decision == ConsentDecision.WITHDRAWN


def build_atlas_kii_pack(*, user):
    """(temp ZIP file, filename, summary): one Word document per interview plus an attribute sheet."""
    content, payloads = _fetch("kii", user)
    form = _Form(content)
    fields = list(_walk(content))
    by_name = {row.get("name"): (section, xpath, row) for section, xpath, row in fields}

    def answer(payload, name):
        if name not in by_name:
            return ""
        _, xpath, row = by_name[name]
        value = payload.get(xpath)
        return "" if value in (None, "") else form.answer_text(row, value)

    entries, attr_rows, skipped_consent, skipped_withdrawn, no_id = {}, [], 0, 0, 0
    attr_header = ["Document", "KII_ID"] + [by_name[n][2].get("name") for n in KII_ATTRIBUTES if n in by_name]
    for payload in sorted(payloads, key=lambda p: (p.get("_submission_time") or "", p.get("_id") or 0)):
        kii_id = answer(payload, "KII_ID").strip()
        if not kii_id:
            no_id += 1
            continue
        if str(payload.get(by_name["consent"][1]) if "consent" in by_name else "yes") == "no":
            skipped_consent += 1
            continue
        if _kii_withdrawn(kii_id):
            skipped_withdrawn += 1
            continue
        doc_name = f"{_var(kii_id)}.docx"
        blocks: list[tuple[str | None, str]] = [("Title", f"{kii_id}: key informant interview")]
        summary_bits = [f"{lab}: {answer(payload, n)}" for n, lab in (("RESP_CAT", "Respondent category"), ("interview_date", "Date"), ("MODE", "Mode")) if answer(payload, n)]
        if summary_bits:
            blocks.append((None, " | ".join(summary_bits)))
        section_seen = None
        for section, xpath, row in fields:
            name = row.get("name")
            if name in KII_LEFT_OUT or name in KII_ATTRIBUTES or name in ("RESP_CAT_OTHER", "CAPTURE_MODE"):
                continue
            text = answer(payload, name)
            if not text:
                continue
            if section != section_seen:
                blocks.append(("Heading1", section))
                section_seen = section
            blocks.append(("Heading2", f"[{name}] {_label(row)}"))
            blocks.append((None, text))
        entries[f"interviews/{doc_name}"] = docx_bytes(blocks)
        attr_rows.append([doc_name, kii_id] + [answer(payload, n) for n in KII_ATTRIBUTES if n in by_name])

    readme = "\r\n".join([
        "ABF-FST key informant interviews, prepared for ATLAS.ti",
        f"Exported {timezone.localtime():%d %b %Y %H:%M} (Africa/Harare) from KoboToolbox through the ABF-FST portal.",
        "",
        f"Interviews in the pack: {len(attr_rows)}",
        f"Left out: {skipped_consent} without participation consent, {skipped_withdrawn} whose participant later withdrew"
        + (f", {no_id} with no interview id" if no_id else "") + ".",
        "",
        "FILES",
        "  interviews/KII-xxxx.docx     one Word document per interview: each question as a heading, the answer beneath it.",
        "  kii_attributes.xlsx          one row per interview: respondent category, mode, date, duration and so on.",
        "",
        "USING IT IN ATLAS.ti",
        "  1. Add all the .docx files in interviews/ as documents. Each document is named after its interview id.",
        "  2. Import kii_attributes.xlsx as document groups / attributes (the first column, Document, matches each",
        "     document's name). The menu wording differs by ATLAS.ti version.",
        "  3. Interview recordings and full transcripts are not held by the portal. Add your transcripts as documents",
        "     named with the same interview id so they share these attributes.",
        "",
        "LEFT OUT ON PURPOSE",
        "  Organisation, job title, interviewer code, consent answers, contact preferences and follow-up referrals.",
        "  Answers are the interviewer's notes as entered; check them for any name before sharing outside the research team.",
        "",
        "Keep ATLAS.ti's AI features switched off for these documents unless the ethics position has been updated to allow it.",
        "Confidential research data. Store and share only as the approved data management plan allows.",
    ])
    entries["kii_attributes.xlsx"] = _xlsx_bytes({"attributes": (attr_header, attr_rows)})
    entries["README.txt"] = readme
    summary = {"pack": "atlas-kii", "interviews": len(attr_rows), "no_consent_left_out": skipped_consent, "withdrawn_left_out": skipped_withdrawn}
    return _zip_response_file(entries), f"ABF-FST_ATLASti_KII_{_stamp()}.zip", summary


# ------------------------------------------------------------------------------ ATLAS.ti: documents

def build_atlas_documents_pack(*, user):
    """(temp ZIP file, filename, summary): the uploaded source files named by DOC-ID plus an attribute sheet."""
    import os

    from apps.evidence.models import DocumentRecord
    from apps.evidence.services import DocumentFileError, open_source_file

    tmp = tempfile.TemporaryFile()
    header = ["Document", "DOC_ID", "Title", "Type", "Publication date", "Author or speaker", "Geographic scope", "Value chain",
              "Source", "Authenticity", "QA status", "Construct tags"]
    rows, notes, missing = [], [], []
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as z:
        for doc in DocumentRecord.objects.exclude(source_file_ref="").order_by("document_id"):
            try:
                path, name, _ = open_source_file(doc)
            except DocumentFileError:
                missing.append(doc.document_id)
                continue
            ext = os.path.splitext(name or path)[1].lower()
            file_name = f"{doc.document_id}_{slugify(doc.title)[:60] or 'document'}{ext}"
            z.write(path, f"documents/{file_name}")
            rows.append([file_name, doc.document_id, doc.title, doc.get_document_type_display(),
                         doc.publication_or_event_date.isoformat() if doc.publication_or_event_date else "",
                         doc.author_or_speaker, doc.geographic_scope, doc.value_chain, doc.source_url_or_reference,
                         doc.get_authenticity_assessment_display(), doc.get_qa_status_display(), "; ".join(doc.construct_tags)])
            if doc.evidence_extract or doc.interpretive_memo:
                notes.append([file_name, doc.document_id, doc.evidence_extract, doc.interpretive_memo])
        z.writestr("documents_attributes.xlsx", _xlsx_bytes({
            "attributes": (header, rows),
            "notes": (["Document", "DOC_ID", "Evidence extract", "Interpretive memo"], notes),
        }))
        z.writestr("README.txt", "\r\n".join([
            "ABF-FST documentary evidence, prepared for ATLAS.ti",
            f"Exported {timezone.localtime():%d %b %Y %H:%M} (Africa/Harare) from the ABF-FST portal.",
            "",
            f"Documents in the pack: {len(rows)} (only records with an uploaded file are included)."
            + (f" Files missing on the server for: {', '.join(missing)}." if missing else ""),
            "",
            "FILES",
            "  documents/                   the uploaded source files, named DOC-ID_short-title.",
            "  documents_attributes.xlsx    one row per document (type, date, author, scope, value chain, authenticity, QA status,",
            "                               construct tags); the notes sheet holds evidence extracts and memos.",
            "",
            "USING IT IN ATLAS.ti",
            "  1. Add every file in documents/ as a document.",
            "  2. Import documents_attributes.xlsx as document groups / attributes (the first column, Document, matches the file name).",
            "",
            "The coded answers for each document are in the Document Analysis Tool workbook on the Data Export screen.",
            "Confidential research data. Store and share only as the approved data management plan allows.",
        ]))
    tmp.seek(0)
    return tmp, f"ABF-FST_ATLASti_Documents_{_stamp()}.zip", {"pack": "atlas-documents", "documents": len(rows), "missing_files": len(missing)}


PACKS = {"smartpls": build_smartpls_pack, "atlas-kii": build_atlas_kii_pack, "atlas-documents": build_atlas_documents_pack}


def audit_export(user, summary: dict) -> None:
    log_action("export.analysis_pack", _Ref(summary["pack"]), summary, user=user)
