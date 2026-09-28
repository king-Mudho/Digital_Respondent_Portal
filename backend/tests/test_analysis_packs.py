"""apps/kobo/analysis_packs.py -- the data as SmartPLS 4 and ATLAS.ti open it. KoboToolbox itself is faked; what is under
test is the shape of the files: numeric-only data with a missing code, nobody who withdrew or failed QA, nothing that
identifies a person in the analysis files, valid Word documents, and that only the PI can download any of it."""

import csv
import io
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.audit.models import AuditEvent
from apps.evidence.models import DocumentRecord
from apps.kobo import analysis_packs as packs

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _row(kind, name, label="", xpath=None):
    return {"type": kind, "name": name, "label": [label or name], "$xpath": xpath or name}


QUESTIONNAIRE = {
    "survey": [
        {"type": "hidden", "name": "sample_id", "$xpath": "sample_id"},
        {"type": "begin_group", "name": "profile", "label": ["Profile"]},
        _row("select_one org_type", "P1", "Organisation type", "profile/P1"),
        _row("select_one yes_no", "E1", "Do you export?", "profile/E1"),
        {"type": "end_group"},
        {"type": "begin_group", "name": "scales", "label": ["Scales"]},
        _row("select_one likert5", "L1", "First scale item", "scales/L1"),
        _row("select_one likert5", "L2", "Second scale item", "scales/L2"),
        _row("select_one strength_na", "S1", "How strong?", "scales/S1"),
        _row("integer", "AGE", "Years trading", "scales/AGE"),
        _row("text", "NOTE1", "Anything else?", "scales/NOTE1"),
        _row("select_multiple finance_source", "FIN", "Finance sources", "scales/FIN"),
        {"type": "end_group"},
    ],
    "choices": (
        [{"list_name": "org_type", "name": n, "label": [n.title()]} for n in ("agribusiness", "farmer", "bank")]
        + [{"list_name": "yes_no", "name": n, "label": [n.title()]} for n in ("yes", "no")]
        + [{"list_name": "likert5", "name": str(n), "label": [str(n)]} for n in range(1, 6)]
        + [{"list_name": "strength_na", "name": n, "label": [n]} for n in ("strong", "moderate", "limited", "none", "not_sure_na")]
        + [{"list_name": "finance_source", "name": n, "label": [n]} for n in ("internal", "bank")]
    ),
}


def _sub(i, **answers):
    p = {"_id": i, "_submission_time": f"2026-10-0{i}T08:00:00", "sample_id": f"SID-2026-00000{i}"}
    p.update(answers)
    return p


PAYLOADS = [
    _sub(1, **{"profile/P1": "farmer", "profile/E1": "yes", "scales/L1": "4", "scales/L2": "5", "scales/S1": "moderate", "scales/AGE": "12", "scales/NOTE1": "free text", "scales/FIN": "internal bank"}),
    _sub(2, **{"profile/P1": "bank", "profile/E1": "no", "scales/L1": "2", "scales/S1": "not_sure_na", "scales/FIN": "bank"}),
    _sub(3, **{"profile/P1": "farmer", "profile/E1": "yes", "scales/L1": "5", "scales/L2": "5"}),     # withdrew later
    _sub(4, **{"profile/P1": "agribusiness", "profile/E1": "no", "scales/L1": "1", "scales/L2": "1"}),  # not through QA yet
]
PORTAL = {1: ["SID-2026-000001", "QA_PASSED", False, "S10"], 2: ["SID-2026-000002", "QA_PASSED", False, "S10"],
          3: ["SID-2026-000003", "QA_PASSED", True, "S10"], 4: ["SID-2026-000004", "QUERY", False, "S09"]}


@pytest.fixture
def smartpls_zip(monkeypatch, db):
    monkeypatch.setattr(packs, "_fetch", lambda key, user: (QUESTIONNAIRE, PAYLOADS))
    monkeypatch.setattr(packs, "_portal_columns", lambda payloads: PORTAL)
    handle, filename, summary = packs.build_smartpls_pack(user=None)
    return zipfile.ZipFile(handle), filename, summary


def _csv(z, name):
    return list(csv.reader(io.StringIO(z.read(name).decode("utf-8"))))


def test_smartpls_file_has_only_accepted_respondents_and_no_one_who_withdrew(smartpls_zip):
    z, _, summary = smartpls_zip
    rows = _csv(z, "data.csv")
    assert len(rows) == 3  # header + respondents 1 and 2
    assert summary["respondents"] == 2 and summary["withdrawn_left_out"] == 1 and summary["not_passed_left_out"] == 1


def test_smartpls_every_data_cell_is_a_number(smartpls_zip):
    z, _, _ = smartpls_zip
    rows = _csv(z, "data.csv")
    for row in rows[1:]:
        for cell in row:
            float(cell)  # raises on any text, so a stray label or id fails here


def test_smartpls_codes_scales_yes_no_categories_and_missing_correctly(smartpls_zip):
    z, _, _ = smartpls_zip
    rows = _csv(z, "data.csv")
    header = rows[0]
    first, second = ({h: v for h, v in zip(header, r)} for r in rows[1:])
    assert first["L1"] == "4" and first["L2"] == "5"                     # the 1-5 scales keep their numbers
    assert first["E1"] == "1" and second["E1"] == "0"                    # yes / no
    assert first["P1"] == "2" and second["P1"] == "3"                    # farmer, bank: numbered in the form's order
    assert first["S1"] == "2" and second["S1"] == str(packs.MISSING)     # "not sure" is missing, not a category
    assert second["L2"] == str(packs.MISSING) and second["AGE"] == str(packs.MISSING)   # unanswered
    assert first["AGE"] == "12"
    assert (first["FIN_internal"], first["FIN_bank"], second["FIN_internal"], second["FIN_bank"]) == ("1", "1", "0", "1")


def test_smartpls_file_holds_no_ids_free_text_or_contact_data(smartpls_zip):
    z, _, _ = smartpls_zip
    header = _csv(z, "data.csv")[0]
    assert "NOTE1" not in header and "sample_id" not in header
    assert header[0] == "respondent_no"
    assert all(v.replace("_", "").isalnum() for v in header)            # variable names SmartPLS can read
    text = z.read("data.csv").decode()
    assert "SID-" not in text and "free text" not in text


def test_smartpls_dictionary_explains_every_variable_and_the_missing_rule(smartpls_zip):
    z, _, _ = smartpls_zip
    rows = _csv(z, "data_dictionary.csv")
    by_var = {r[0]: r for r in rows[1:]}
    assert set(by_var) == set(_csv(z, "data.csv")[0][1:])
    assert "2=Farmer" in by_var["P1"][3] and "1=Yes" in by_var["E1"][3]
    assert str(packs.MISSING) in by_var["S1"][4]
    assert "-99" in z.read("README.txt").decode()


def test_smartpls_xlsx_matches_the_csv_and_the_key_stays_in_the_private_folder(smartpls_zip):
    z, _, _ = smartpls_zip
    sheet = load_workbook(io.BytesIO(z.read("data.xlsx")))["data"]
    assert [[str(c) for c in r] for r in sheet.iter_rows(values_only=True)] == _csv(z, "data.csv")
    key = _csv(z, "PRIVATE_do_not_share/respondent_key.csv")
    assert key[1][:2] == ["1", "SID-2026-000001"] and len(key) == 3
    assert not any(n.endswith("respondent_key.csv") and "PRIVATE" not in n for n in z.namelist())


def test_smartpls_can_include_submissions_not_yet_through_qa_but_never_withdrawn_ones(monkeypatch, db):
    monkeypatch.setattr(packs, "_fetch", lambda key, user: (QUESTIONNAIRE, PAYLOADS))
    monkeypatch.setattr(packs, "_portal_columns", lambda payloads: PORTAL)
    handle, _, summary = packs.build_smartpls_pack(user=None, include_unreviewed=True)
    assert summary["respondents"] == 3 and summary["withdrawn_left_out"] == 1


# ------------------------------------------------------------------------------------------------- ATLAS.ti KII

KII_FORM = {
    "survey": [
        {"type": "begin_group", "name": "part_a", "label": ["Part A"]},
        _row("text", "KII_ID", "Interview ID", "part_a/KII_ID"),
        _row("date", "interview_date", "Date", "part_a/interview_date"),
        _row("text", "org_represented", "Organisation represented", "part_a/org_represented"),
        _row("text", "official_title", "Official title", "part_a/official_title"),
        _row("select_one respondent_category", "RESP_CAT", "Respondent category", "part_a/RESP_CAT"),
        _row("select_one yes_no", "consent", "Consent to participate", "part_a/consent"),
        _row("integer", "duration_minutes", "Duration", "part_a/duration_minutes"),
        {"type": "end_group"},
        {"type": "begin_group", "name": "part_b", "label": ["Part B: Core questions"]},
        _row("text", "K1", "What makes an agribusiness bankable?", "part_b/K1"),
        _row("text", "K2", "What are the barriers?", "part_b/K2"),
        {"type": "end_group"},
        {"type": "begin_group", "name": "part_e", "label": ["Part E: Close"]},
        _row("text", "followup_referral", "Who else should we speak to?", "part_e/followup_referral"),
        _row("text", "analytical_memo", "Interviewer memo", "part_e/analytical_memo"),
        {"type": "end_group"},
    ],
    "choices": [{"list_name": "respondent_category", "name": "bank_mfi", "label": ["Bank/MFI"]},
                {"list_name": "yes_no", "name": "yes", "label": ["Yes"]}, {"list_name": "yes_no", "name": "no", "label": ["No"]}],
}


def _kii(i, **extra):
    p = {"_id": i, "_submission_time": f"2026-10-1{i}T09:00:00", "part_a/KII_ID": f"KII-000{i}", "part_a/interview_date": "2026-10-10",
         "part_a/org_represented": "Secret Bank Ltd", "part_a/official_title": "Chief Executive", "part_a/RESP_CAT": "bank_mfi",
         "part_a/consent": "yes", "part_a/duration_minutes": "45", "part_b/K1": "Cash flow discipline.\nAnd records.",
         "part_b/K2": "Collateral <rules> & tenure.", "part_e/followup_referral": "Call Mr X on 077...", "part_e/analytical_memo": "Rich interview."}
    p.update(extra)
    return p


@pytest.fixture
def kii_zip(monkeypatch, db):
    payloads = [_kii(1), _kii(2, **{"part_a/consent": "no"}), _kii(3), _kii(4)]
    monkeypatch.setattr(packs, "_fetch", lambda key, user: (KII_FORM, payloads))
    monkeypatch.setattr(packs, "_kii_withdrawn", lambda kii_id: kii_id == "KII-0004")
    handle, filename, summary = packs.build_atlas_kii_pack(user=None)
    return zipfile.ZipFile(handle), summary


def _docx_text(z, name):
    inner = zipfile.ZipFile(io.BytesIO(z.read(name)))
    root = ET.fromstring(inner.read("word/document.xml"))
    return [(p.find(f"{W}pPr/{W}pStyle").get(f"{W}val") if p.find(f"{W}pPr/{W}pStyle") is not None else None,
             "".join(t.text or "" for t in p.iter(f"{W}t"))) for p in root.iter(f"{W}p")]


def test_kii_pack_has_one_valid_word_document_per_consenting_interview(kii_zip):
    z, summary = kii_zip
    docs = sorted(n for n in z.namelist() if n.startswith("interviews/"))
    assert docs == ["interviews/KII_0001.docx", "interviews/KII_0003.docx"]
    assert summary["no_consent_left_out"] == 1 and summary["withdrawn_left_out"] == 1
    paras = _docx_text(z, "interviews/KII_0001.docx")
    assert ("Title", "KII-0001: key informant interview") in paras


def test_kii_document_has_each_question_as_a_heading_with_the_answer_under_it(kii_zip):
    z, _ = kii_zip
    paras = _docx_text(z, "interviews/KII_0001.docx")
    texts = [t for _, t in paras]
    i = texts.index("[K1] What makes an agribusiness bankable?")
    assert paras[i][0] == "Heading2" and texts[i + 1:i + 3] == ["Cash flow discipline.", "And records."]
    assert "Collateral <rules> & tenure." in texts                         # special characters survive
    assert ("Heading1", "Part B: Core questions") in paras
    assert "Interviewer memo" in " ".join(texts)


def test_kii_documents_and_attributes_hold_nothing_that_names_the_person(kii_zip):
    z, _ = kii_zip
    every = " ".join(t for n in z.namelist() if n.endswith(".docx") for _, t in _docx_text(z, n))
    for hidden in ("Secret Bank", "Chief Executive", "Mr X", "077"):
        assert hidden not in every
    sheet = load_workbook(io.BytesIO(z.read("kii_attributes.xlsx")))["attributes"]
    values = [c for r in sheet.iter_rows(values_only=True) for c in r]
    assert "Secret Bank Ltd" not in values and "Chief Executive" not in values
    header = [c.value for c in sheet[1]]
    assert header[:2] == ["Document", "KII_ID"] and "RESP_CAT" in header
    row = {h: c.value for h, c in zip(header, sheet[2])}
    assert row["Document"] == "KII_0001.docx" and row["RESP_CAT"] == "Bank/MFI" and str(row["duration_minutes"]) == "45"


def test_word_document_opens_as_a_real_docx(tmp_path):
    inner = zipfile.ZipFile(io.BytesIO(packs.docx_bytes([("Title", "T"), ("Heading2", "Q"), (None, "a\nb")])))
    assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml"} <= set(inner.namelist())
    for name in inner.namelist():
        ET.fromstring(inner.read(name))  # every part is well-formed XML


# ----------------------------------------------------------------------------------- ATLAS.ti documents + view

def test_documents_pack_names_files_by_doc_id_and_lists_their_attributes(settings, tmp_path, db):
    settings.PRIVATE_DATA_ROOT = tmp_path
    (tmp_path / "documents").mkdir()
    (tmp_path / "documents" / "abc.pdf").write_bytes(b"%PDF-1.4 fake")
    DocumentRecord.objects.create(document_id="DOC-0007", title="Vision 2030: the plan", document_type="OFFICIAL", value_chain="Cereals",
                                  source_file_ref="documents/abc.pdf", source_file_name="Vision.PDF", interpretive_memo="A memo.")
    DocumentRecord.objects.create(document_id="DOC-0008", title="No file yet", document_type="OFFICIAL")
    handle, _, summary = packs.build_atlas_documents_pack(user=None)
    z = zipfile.ZipFile(handle)
    assert "documents/DOC-0007_vision-2030-the-plan.pdf" in z.namelist()
    assert z.read("documents/DOC-0007_vision-2030-the-plan.pdf") == b"%PDF-1.4 fake"
    book = load_workbook(io.BytesIO(z.read("documents_attributes.xlsx")))
    assert [c.value for c in book["attributes"][2]][:3] == ["DOC-0007_vision-2030-the-plan.pdf", "DOC-0007", "Vision 2030: the plan"]
    assert book["notes"][2][3].value == "A memo."
    assert summary == {"pack": "atlas-documents", "documents": 1, "missing_files": 0}


def _client(role_name):
    role, _ = Role.objects.get_or_create(name=role_name)
    user = User.objects.create_user(username=f"pack_{role_name.lower()}", password="x", role=role)
    client = APIClient()
    client.force_authenticate(user)
    return client, user


@pytest.mark.parametrize("pack", ["smartpls", "atlas-kii", "atlas-documents"])
def test_only_the_pi_can_download_a_pack(pack, db):
    for role in (Role.ANALYST, Role.FIELD_COORDINATOR, Role.KII_RA, Role.SUPERVISOR_READONLY):
        client, _ = _client(role)
        assert client.get(f"/api/v1/kobo/analysis-packs/{pack}/").status_code == 403


def test_the_pi_downloads_a_pack_and_the_download_is_audited(monkeypatch, db):
    monkeypatch.setattr(packs, "_fetch", lambda key, user: (QUESTIONNAIRE, PAYLOADS))
    monkeypatch.setattr(packs, "_portal_columns", lambda payloads: PORTAL)
    client, user = _client(Role.PI_ADMIN)
    response = client.get("/api/v1/kobo/analysis-packs/smartpls/")
    assert response.status_code == 200 and response["Content-Type"] == "application/zip"
    assert "attachment" in response["Content-Disposition"] and response["Cache-Control"] == "no-store"
    assert "data.csv" in zipfile.ZipFile(io.BytesIO(b"".join(response.streaming_content))).namelist()
    event = AuditEvent.objects.get(action="export.analysis_pack")
    assert event.user_id == user.id and event.metadata["respondents"] == 2


def test_an_unknown_pack_is_refused(db):
    client, _ = _client(Role.PI_ADMIN)
    assert client.get("/api/v1/kobo/analysis-packs/spss/").status_code == 404
