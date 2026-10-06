"""
Respondents and outreach workbook (apps/dashboards/outreach_export.py). Under test: one row per person on every
contactable case (and a row for a case with nobody yet), a locked Reserve left out, each case's latest invitation with
who issued it and whether the portal emailed it, reminders and contact attempts counted, KII informants and every
invitation listed -- and that only the PI can download it, with each download audited and no contact detail in the
audit entry.
"""

from io import BytesIO

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.audit.models import AuditEvent
from apps.audit.utils import log_action
from apps.contacts.models import ContactEvent, Respondent
from apps.invitations.services import issue_invitation
from apps.kii.services import create_kii_record, issue_kii_invitation
from apps.messaging.services import record_manual_follow_up
from apps.sampling.models import ReserveStatus, SampleType
from apps.sampling.services import create_sample_case

URL = "/api/v1/export/outreach/"


def _user(role_name, username):
    role, _ = Role.objects.get_or_create(name=role_name)
    return User.objects.create_user(username=username, password="x", role=role)


def _client(user):
    client = APIClient()
    client.force_authenticate(user)
    return client


def _sheets(response):
    wb = load_workbook(BytesIO(response.content))
    out = {}
    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        header = rows[0]
        out[sheet.title] = [dict(zip(header, row)) for row in rows[1:]] if sheet.title != "About this file" else rows
    return out


@pytest.fixture
def outreach(main_case, organisation, stratum, locked_reserve_case):
    pi = _user(Role.PI_ADMIN, "out_pi")
    Respondent.objects.create(sample_case=main_case, full_name="Jane Doe", is_eligible=True, phone="0773943709",
                              whatsapp_number="0773943709", email="jane@testorg.co.zw")
    Respondent.objects.create(sample_case=main_case, full_name="Peter Moyo", phone="0242700000")
    for status in ("S01", "S02", "S03"):
        from apps.sampling.services import transition_workflow_status

        transition_workflow_status(main_case, status)
    issue_invitation(main_case, channel="WHATSAPP", issued_by=pi)
    _, _, emailed_token = issue_invitation(main_case, channel="EMAIL", issued_by=pi)
    log_action("invitation.emailed", emailed_token, {"sent_to": "j***@testorg.co.zw"}, user=pi)
    record_manual_follow_up(sample_case=main_case, template_name="drp_reminder_day2", user=pi)
    ContactEvent.objects.create(sample_case=main_case, channel="PHONE", outcome="NO_ANSWER", occurred_at=timezone.now(), ra=pi)

    empty_case = create_sample_case(organisation=organisation, stratum=stratum, sample_type=SampleType.MAIN, year=2026)
    activated = create_sample_case(organisation=organisation, stratum=stratum, sample_type=SampleType.RESERVE, year=2026)
    activated.status = ReserveStatus.ACTIVATED
    activated.save(update_fields=["status"])

    kii = create_kii_record(participant_name="Tendai Moyo", participant_role="Head of Agribusiness",
                            stakeholder_category="Bank, DFI & MFI", email="tmoyo@cbz.co.zw",
                            metadata={"organisation_name": "CBZ Bank"})
    issue_kii_invitation(kii, channel="EMAIL", issued_by=pi)
    return {"pi": pi, "empty": empty_case, "activated": activated, "locked": locked_reserve_case, "kii": kii}


def test_the_pi_downloads_a_workbook_with_every_contactable_person_kii_informant_and_invitation(main_case, outreach):
    response = _client(outreach["pi"]).get(URL)
    assert response.status_code == 200
    assert response["Content-Type"].startswith("application/vnd.openxmlformats")
    assert "respondents_and_outreach" in response["Content-Disposition"]
    sheets = _sheets(response)
    assert list(sheets) == ["Main-400 respondents", "KII informants", "Invitation log", "About this file"]

    main = sheets["Main-400 respondents"]
    mine = [row for row in main if row["Sample ID"] == main_case.sample_id]
    assert [row["Respondent"] for row in mine] == ["Jane Doe", "Peter Moyo"]  # eligible first, one row each
    jane = mine[0]
    assert (jane["Phone"], jane["Email"], jane["Eligible"]) == ("0773943709", "jane@testorg.co.zw", "Yes")
    assert jane["Invitations issued"] == 2 and jane["Latest invitation: channel"] == "EMAIL"
    assert jane["Latest invitation: by"] == "out_pi" and jane["Emailed by the portal"]
    assert jane["Reminders recorded as sent"] == 1 and jane["Contact attempts logged"] == 1
    assert jane["Last attempt: channel / outcome"] == "PHONE / NO_ANSWER"

    ids = {row["Sample ID"] for row in main}
    assert outreach["empty"].sample_id in ids  # a case with nobody recorded still has a row
    assert outreach["activated"].sample_id in ids and outreach["locked"].sample_id not in ids

    kii = sheets["KII informants"][0]
    assert (kii["KII ID"], kii["Organisation"], kii["Email"], kii["Invitations issued"]) == (
        outreach["kii"].kii_id, "CBZ Bank", "tmoyo@cbz.co.zw", 1)
    log = sheets["Invitation log"]
    assert sorted(row["Type"] for row in log) == ["KII", "Main-400", "Main-400"]
    assert {row["Channel"] for row in log if row["Type"] == "Main-400"} == {"WHATSAPP", "EMAIL"}


@pytest.mark.parametrize("role", [Role.FIELD_COORDINATOR, Role.SUPERVISOR_READONLY, Role.ANALYST, Role.CONTACT_RA, Role.KII_RA])
def test_only_the_pi_can_download_it(outreach, role):
    assert _client(_user(role, f"out_{role.lower()}")).get(URL).status_code == 403
    assert APIClient().get(URL).status_code in (401, 403)


def test_each_download_is_audited_with_counts_and_no_contact_details(outreach):
    _client(outreach["pi"]).get(URL)
    event = AuditEvent.objects.get(action="export.outreach_generated")
    assert event.user == outreach["pi"]
    assert event.metadata["kii_records"] == 1 and event.metadata["invitations"] == 3 and event.metadata["people"] == 2
    assert "0773943709" not in str(event.metadata) and "jane@" not in str(event.metadata)
