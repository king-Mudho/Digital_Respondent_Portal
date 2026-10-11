"""Importing Main-400 background research from a ChatGPT workbook (apps/proit/research_import.py, target "main",
2026-10-11). The Main-400 pack has a Sample ID column, Fill A Contacts and Fill B PROIT facts, and no interview-prep
sheet. Under test: contacts become suggestions on the case and facts become PROIT proposals, nothing lands on the case
itself; accepting an imported person creates the respondent through the usual flow; a locked Reserve and unknown cases
are left out; interview-prep notes are refused for Main-400; only the PI and FC may import."""

from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.audit.models import AuditEvent
from apps.contacts.models import ContactProposal, Respondent
from apps.proit.models import AIProposal, PreProfile

URL = "/api/v1/proit/main-import/"
A = ["Sample ID", "Organisation", "Website", "Main phone", "General email", "Physical address",
     "Owner / manager (official source only)", "Their title", "Their published work phone", "Their published work email",
     "Source links", "Confidence", "Notes"]
B = ["Sample ID", "Organisation", "field_id (sheet 4)", "Field label", "Value found (documentary)", "Confidence", "Source title",
     "Publisher", "Source date", "Link", "Source tier (sheet 5)", "Conflicting source?", "Notes"]


def _user(role, username):
    role_obj, _ = Role.objects.get_or_create(name=role)
    return User.objects.create_user(username=username, password="x", role=role_obj)


def _workbook(rows_a=(), rows_b=(), with_prep=False) -> bytes:
    wb = Workbook()
    for title, header, rows in (("Fill A Contacts", A, rows_a), ("Fill B PROIT facts", B, rows_b)):
        ws = wb.create_sheet(title)
        ws.append(header)
        ws.append(["SID-2026-000000", "Example Agro EXAMPLE (delete)"] + ["example"] * (len(header) - 2))
        for row in rows:
            ws.append(row)
    if with_prep:
        ws = wb.create_sheet("Fill C Interview prep")
        ws.append(["Sample ID", "Organisation", "Known evidence summary"])
    del wb["Sheet"]
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _post(user, data, *, dry_run):
    client = APIClient()
    client.force_authenticate(user)
    return client.post(URL, {"file": SimpleUploadedFile("main.xlsx", data), "dry_run": "true" if dry_run else "false"}, format="multipart")


def _rows(sample_id):
    return ([[sample_id, "Test Org", "https://testorg.co.zw", "0242 700 000", "info@testorg.co.zw", "1 Farm Road, Harare",
              "Rudo Chikore", "Managing Director", "0773 943 709", "", "https://testorg.co.zw/about", "MODERATE", ""]],
            [[sample_id, "Test Org", "year_established", "Year established", "2011", "HIGH", "PRAZ register", "PRAZ", "2025",
              "https://praz.org.zw/x", "Tier 1", "No", ""]])


@pytest.fixture
def pi(db):
    return _user(Role.PI_ADMIN, "mimp_pi")


def test_contacts_and_facts_become_suggestions_on_the_case(pi, main_case):
    a, b = _rows(main_case.sample_id)
    resp = _post(pi, _workbook(a, b), dry_run=False)
    assert resp.status_code == 201
    [row] = resp.json()["records"]
    assert (row["record"], row["record_id"], row["status"], row["contacts"], row["facts"]) == (
        main_case.sample_id, main_case.sample_id, "imported", 5, 1)
    assert set(ContactProposal.objects.filter(sample_case=main_case).values_list("kind", flat=True)) == {
        "ORG_PHONE", "ORG_EMAIL", "WEBSITE", "OFFICE_LOCATION", "PERSON"}
    assert not Respondent.objects.filter(sample_case=main_case).exists()  # nothing on the case until accepted
    profile = PreProfile.objects.get(sample_case=main_case)
    assert AIProposal.objects.get(pre_profile=profile).proposed_value == "2011"
    assert AuditEvent.objects.get(action="proit.main_research_imported").metadata["contacts"] == 5
    assert resp.json()["missing_sheets"] == []  # Main-400 packs have no Fill C, and that is not reported as missing


def test_accepting_an_imported_person_creates_the_respondent_through_the_usual_flow(pi, main_case):
    from apps.contacts.contact_finder import accept_proposal

    a, b = _rows(main_case.sample_id)
    _post(pi, _workbook(a, b), dry_run=False)
    accept_proposal(ContactProposal.objects.get(sample_case=main_case, kind="PERSON"), user=pi)
    person = Respondent.objects.get(sample_case=main_case)
    assert person.full_name == "Rudo Chikore" and "773943709" in person.phone.replace(" ", "")


def test_a_locked_reserve_and_unknown_cases_are_left_out(pi, locked_reserve_case):
    a, _ = _rows(locked_reserve_case.sample_id)
    unknown, _ = _rows("SID-2026-999999")
    resp = _post(pi, _workbook(a + unknown), dry_run=False)
    rows = {r["record"]: r for r in resp.json()["records"]}
    assert rows[locked_reserve_case.sample_id]["status"] == "skipped" and "locked Reserve" in rows[locked_reserve_case.sample_id]["warnings"][0]
    assert rows["SID-2026-999999"]["status"] == "skipped"
    assert not ContactProposal.objects.exists()


def test_sample_ids_are_matched_however_they_are_typed(pi, main_case):
    loose = main_case.sample_id.replace("SID-", "sid ").replace("-", " ")
    a, _ = _rows(loose)
    resp = _post(pi, _workbook(a), dry_run=True)
    assert resp.json()["records"][0]["record"] == main_case.sample_id and resp.json()["records"][0]["status"] == "ready"


def test_interview_prep_is_for_kii_only(pi, main_case):
    from openpyxl import load_workbook

    a, b = _rows(main_case.sample_id)
    wb = load_workbook(BytesIO(_workbook(a, b, with_prep=True)))
    wb["Fill C Interview prep"].append([main_case.sample_id, "Test Org", "Something known"])
    out = BytesIO()
    wb.save(out)
    resp = _post(pi, out.getvalue(), dry_run=False)
    assert "KII records only" in " ".join(resp.json()["records"][0]["warnings"])
    assert PreProfile.objects.get(sample_case=main_case).known_evidence_summary == ""


@pytest.mark.parametrize("role", [Role.CONTACT_RA, Role.SUPERVISOR_READONLY, Role.KII_RA])
def test_only_the_pi_and_field_coordinator_import(main_case, role):
    a, b = _rows(main_case.sample_id)
    assert _post(_user(role, f"mimp_{role.lower()}"), _workbook(a, b), dry_run=True).status_code == 403
    assert _post(_user(Role.FIELD_COORDINATOR, "mimp_fc"), _workbook(a, b), dry_run=True).status_code == 200
