"""PI's full KoboToolbox data exports: workbook and PDF ZIP (added 2026-09-15)."""

import io
import zipfile
from unittest.mock import patch

import pytest
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import Role, User
from apps.audit.models import AuditEvent
from apps.kobo.models import QAStatus, QUANSubmission

CONTENT = {
    "survey": [
        {"type": "start", "name": "start", "$xpath": "start"},
        {"type": "hidden", "name": "sample_id", "$xpath": "sample_id"},
        {"type": "begin_group", "name": "profile", "label": [None, "Profile"], "$xpath": "profile"},
        {"type": "select_one", "select_from_list_name": "likert5", "name": "NFM1",
         "label": [None, "Buyer contracts improve access"], "$xpath": "profile/NFM1"},
        {"type": "select_one", "select_from_list_name": "org_type", "name": "P1",
         "label": [None, "Organisation type"], "$xpath": "profile/P1"},
        {"type": "text", "name": "D3", "label": [None, "Biggest barrier"], "$xpath": "profile/D3"},
        {"type": "note", "name": "n", "label": [None, "A note"], "$xpath": "profile/n"},
        {"type": "end_group", "name": "profile"},
        {"type": "begin_repeat", "name": "metric_repeat", "label": [None, "Metric"], "$xpath": "metric_repeat"},
        {"type": "integer", "name": "value", "label": [None, "Value"], "$xpath": "metric_repeat/value"},
        {"type": "end_repeat", "name": "metric_repeat"},
    ],
    "choices": [
        {"list_name": "likert5", "name": "4", "label": [None, "Agree"]},
        {"list_name": "org_type", "name": "processor", "label": [None, "Processor"]},
    ],
}


def _payloads(sample_id):
    return [{
        "_id": 11, "_uuid": "u-11", "meta/rootUuid": "uuid:u-11", "_submission_time": "2026-09-15T08:00:00",
        "sample_id": sample_id, "SAMPLE_ID_FINAL": sample_id,
        "profile/NFM1": "4", "profile/P1": "processor", "profile/D3": "=HYPERLINK(\"http://x\")",
        "metric_repeat": [{"metric_repeat/value": "120"}, {"metric_repeat/value": "7"}],
    }]


@pytest.fixture
def kobo(settings, main_case):
    settings.KOBO_API_TOKEN, settings.KOBO_ASSET_UID = "token", "q-asset"
    with patch("apps.kobo.submission_copies._kobo_get", return_value={"content": CONTENT}), \
         patch("apps.kobo.data_export.KoboClient.fetch_submissions", return_value=_payloads(main_case.sample_id)):
        yield main_case


def _client(role_name, username):
    role, _ = Role.objects.get_or_create(name=role_name)
    client = APIClient()
    client.force_authenticate(User.objects.create_user(username=username, password="x", role=role))
    return client


@pytest.mark.django_db
def test_the_workbook_has_codes_labels_dictionary_repeats_and_portal_columns(kobo):
    QUANSubmission.objects.create(sample_case=kobo, kobo_submission_uuid="u-11", administration_mode="01",
                                  submitted_at=timezone.now(), qa_status=QAStatus.QA_PASSED)
    resp = _client(Role.PI_ADMIN, "dx_pi").get("/api/v1/kobo/forms/questionnaire/export/xlsx/")
    assert resp.status_code == 200 and resp["Content-Disposition"].endswith('.xlsx"')

    wb = load_workbook(io.BytesIO(resp.content))
    assert {"README", "data_codes", "data_labels", "questions", "choices", "repeat_1"} <= set(wb.sheetnames)

    codes = list(wb["data_codes"].values)
    header, row = codes[0], dict(zip(codes[0], codes[1]))
    assert "profile/n" not in header  # notes are not data
    assert row["profile/NFM1"] == 4  # a number, ready for analysis
    assert row["profile/D3"] == '=HYPERLINK("http://x")' and wb["data_codes"].cell(2, header.index("profile/D3") + 1).data_type == "s"
    assert (row["portal_matched_case"], row["portal_qa_status"]) == (kobo.sample_id, "QA_PASSED")

    labels = dict(zip(*list(wb["data_labels"].values)[:2]))
    assert (labels["profile/NFM1"], labels["profile/P1"]) == ("Agree", "Processor")
    assert [r[:3] for r in list(wb["repeat_1"].values)[1:]] == [(11, 1, 120), (11, 2, 7)]
    assert ("profile/NFM1", "Buyer contracts improve access") in [r[:2] for r in wb["questions"].values]
    assert AuditEvent.objects.filter(action="kobo.data_exported").exists()


@pytest.mark.django_db
def test_every_completed_form_downloads_as_pdfs_in_a_zip(kobo):
    resp = _client(Role.PI_ADMIN, "dx_pi2").get("/api/v1/kobo/forms/questionnaire/export/pdfs/")
    assert resp.status_code == 200
    archive = zipfile.ZipFile(io.BytesIO(b"".join(resp.streaming_content)))
    names = archive.namelist()
    assert "manifest.csv" in names
    [pdf] = [n for n in names if n.endswith(".pdf")]
    assert archive.read(pdf).startswith(b"%PDF-")


@pytest.mark.django_db
@pytest.mark.parametrize("role", [Role.FIELD_COORDINATOR, Role.ANALYST, Role.QUAN_QA_RA])
def test_full_data_exports_are_for_the_pi_only(kobo, role):
    client = _client(role, f"dx_{role.lower()}")
    assert client.get("/api/v1/kobo/forms/questionnaire/export/xlsx/").status_code == 403
    assert client.get("/api/v1/kobo/forms/questionnaire/export/pdfs/").status_code == 403
